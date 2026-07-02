"""Parse the Excel registry into normalized catalog records.

Workbook layout (per data tab):
  row 1 : free-text title / metadata
  row 2 : header  -> 一级分类 | 二级分类 | 渠道名称 | 来源链接 | 方式 |
                     是否付费 | 账号 | 频率 | 备注 | 优先级 | (Bio, tab 9 only)
  row 3+: data
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .module_spec import MODULES, ModuleSpec

# Canonical Chinese header -> field name on CatalogRow.
HEADER_MAP = {
    "一级分类": "l1_category",
    "二级分类": "l2_category",
    "渠道名称": "name",
    "来源链接": "url",
    "方式": "method",
    "是否付费": "paid",
    "账号": "account",
    "频率": "frequency",
    "备注": "remark",
    "优先级": "priority",
}

# Leading status glyphs used in the 备注 column.
_HEALTH_RE = re.compile(r"^\s*([🟢🟡🔴⚪✅❌⛔🛠️📊⏳])")


@dataclass
class CatalogRow:
    module: str            # ModuleKey value
    tab: str
    l1_category: str | None = None
    l2_category: str | None = None
    name: str | None = None
    url: str | None = None
    method: str | None = None
    paid: str | None = None
    account: str | None = None
    frequency: str | None = None
    remark: str | None = None
    priority: str | None = None
    health: str | None = None
    extra: dict = field(default_factory=dict)


def _clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _resolve_columns(header_row) -> dict[int, str]:
    """Map 1-based column index -> field name using the header labels."""
    cols: dict[int, str] = {}
    for idx, cell in enumerate(header_row, start=1):
        label = _clean(cell.value)
        if label and label in HEADER_MAP:
            cols[idx] = HEADER_MAP[label]
    return cols


def _extract_health(remark: str | None) -> str | None:
    if not remark:
        return None
    m = _HEALTH_RE.match(remark)
    return m.group(1) if m else None


def load_module(ws, spec: ModuleSpec) -> list[CatalogRow]:
    rows: list[CatalogRow] = []
    header = next(ws.iter_rows(min_row=2, max_row=2))
    cols = _resolve_columns(header)
    if not cols:
        return rows

    # Sheet 9 carries an extra "Bio" column; capture it into extra/payload.
    bio_col = None
    for idx, cell in enumerate(header, start=1):
        label = _clean(cell.value)
        if label and label.lower().startswith("bio"):
            bio_col = idx

    for raw in ws.iter_rows(min_row=3):
        values = {field_name: _clean(raw[idx - 1].value) for idx, field_name in cols.items()}
        if not values.get("name"):
            continue  # skip spacer / footnote rows
        row = CatalogRow(module=spec.key.value, tab=spec.tab, **values)
        row.health = _extract_health(row.remark)
        if bio_col and bio_col - 1 < len(raw):
            bio = _clean(raw[bio_col - 1].value)
            if bio:
                row.extra["bio"] = bio
        rows.append(row)
    return rows


def load_catalog(path: str | Path) -> list[CatalogRow]:
    """Load every data tab into a flat list of normalized catalog rows."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out: list[CatalogRow] = []
    for spec in MODULES:
        if spec.tab in wb.sheetnames:
            out.extend(load_module(wb[spec.tab], spec))
    wb.close()
    return out
